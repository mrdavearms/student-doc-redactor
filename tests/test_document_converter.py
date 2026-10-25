"""
LibreOffice is run with a profile private to this process.

Without -env:UserInstallation, soffice hands the headless job to any
LibreOffice window the user has open. A good document still converts, but a
damaged one then exits 0 with no PDF and no error text at all (reproduced on
LibreOffice 26.8), so the user saw "Conversion failed:" and nothing after it.
"""
import sys, os, shutil
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src', 'core'))

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import document_converter
from document_converter import DocumentConverter


def _converter():
    with patch('binary_resolver.resolve_libreoffice', return_value='/fake/soffice'):
        conv = DocumentConverter()
    conv.soffice_path = '/fake/soffice'
    return conv


def _run_recording(calls, returncode=0, stdout='', stderr=''):
    def fake_run(cmd, **kwargs):
        calls.append(cmd)
        return SimpleNamespace(returncode=returncode, stdout=stdout, stderr=stderr)
    return fake_run


class TestPrivateProfile:

    def test_soffice_gets_a_profile_under_the_temp_root_not_the_users_folder(self, tmp_path):
        calls = []
        docx = tmp_path / "docs" / "Billy Bob report.docx"
        docx.parent.mkdir()
        with patch('document_converter.subprocess.run', _run_recording(calls)):
            _converter().convert_to_pdf(docx, tmp_path / "out")

        cmd = calls[0]
        env_args = [a for a in cmd if a.startswith('-env:UserInstallation=')]
        assert len(env_args) == 1, cmd
        uri = env_args[0].split('=', 1)[1]
        assert uri.startswith('file://')
        profile = Path(uri[len('file://'):])
        assert profile.is_relative_to(document_converter._temp_root())
        assert not profile.is_relative_to(tmp_path)
        # Before the mode flags, where soffice expects bootstrap variables.
        assert cmd.index(env_args[0]) < cmd.index('--headless')


class TestFailureMessages:

    def test_exit_zero_with_no_pdf_and_no_output_still_explains_itself(self, tmp_path):
        docx = tmp_path / "Damaged report.docx"
        with patch('document_converter.subprocess.run', _run_recording([], 0, '', '')):
            ok, message, path = _converter().convert_to_pdf(docx, tmp_path / "out")
        assert not ok
        assert path is None
        assert message.startswith("Conversion failed: ")
        assert len(message) > len("Conversion failed: ")
        assert "no PDF" in message

    def test_libreoffice_error_text_is_passed_through(self, tmp_path):
        docx = tmp_path / "Damaged report.docx"
        with patch('document_converter.subprocess.run',
                   _run_recording([], 1, '', 'Error: source file could not be loaded\n')):
            ok, message, _ = _converter().convert_to_pdf(docx, tmp_path / "out")
        assert not ok
        assert message == "Conversion failed: Error: source file could not be loaded"


class TestTempRootFromElectron:

    def test_temp_root_uses_env_when_set(self, tmp_path, monkeypatch):
        monkeypatch.setattr(document_converter, '_TEMP_ROOT', None)
        monkeypatch.setenv('REDACTION_TEMP_ROOT', str(tmp_path))
        with patch('document_converter.atexit.register') as reg:
            assert document_converter._temp_root() == tmp_path
        reg.assert_called_once()

    def test_temp_root_ignores_relative_env(self, tmp_path, monkeypatch):
        monkeypatch.setattr(document_converter, '_TEMP_ROOT', None)
        monkeypatch.chdir(tmp_path)
        (tmp_path / 'rel').mkdir()
        monkeypatch.setenv('REDACTION_TEMP_ROOT', 'rel')
        with patch('document_converter.atexit.register'):
            root = document_converter._temp_root()
        try:
            assert root.is_absolute() and root != tmp_path / 'rel'
            assert root.name.startswith('redaction-tool-py-')
        finally:
            shutil.rmtree(root, True)


# ---------------------------------------------------------------------------
# Rule 82: LibreOffice opens a third-party Word file with macros, active
# content and remote links switched off. The app's promise is that nothing
# leaves the machine; a hostile .docx must not be able to make the conversion
# fetch a URL or run code.
# ---------------------------------------------------------------------------

import http.server
import re
import threading
import subprocess
import zipfile

import pytest

from binary_resolver import resolve_libreoffice

_LOCKED_VALUES = {
    ('/org.openoffice.Office.Common/Security/Scripting', 'DisableMacrosExecution'): 'true',
    ('/org.openoffice.Office.Common/Security/Scripting', 'MacroSecurityLevel'): '3',
    ('/org.openoffice.Office.Common/Security/Scripting', 'BlockUntrustedRefererLinks'): 'true',
    ('/org.openoffice.Office.Common/Security/Scripting', 'DisableActiveContent'): 'true',
    ('/org.openoffice.Office.Writer/Content/Update', 'Link'): '0',
}


def _xcu_values(xcu_text):
    """{(item path, prop name): value} from a registrymodifications.xcu."""
    found = {}
    for path, name, value in re.findall(
            r'<item oor:path="([^"]+)"><prop oor:name="([^"]+)" oor:op="fuse">'
            r'<value>([^<]*)</value></prop></item>', xcu_text):
        found[(path, name)] = value
    return found


class TestLockedProfile:

    def test_profile_written_before_conversion(self, tmp_path, monkeypatch):
        monkeypatch.setattr(document_converter, '_TEMP_ROOT', tmp_path / 'root')
        (tmp_path / 'root').mkdir()
        seen = {}

        def fake_run(cmd, **kwargs):
            uri = [a for a in cmd if a.startswith('-env:UserInstallation=')][0]
            profile = Path(uri.split('=', 1)[1][len('file://'):])
            xcu = profile / 'user' / 'registrymodifications.xcu'
            seen['text'] = xcu.read_text(encoding='utf-8') if xcu.exists() else None
            return SimpleNamespace(returncode=1, stdout='', stderr='x')

        with patch('document_converter.subprocess.run', fake_run):
            _converter().convert_to_pdf(tmp_path / 'Billy Bob.docx', tmp_path / 'out')

        assert seen['text'] is not None, "profile was not written before soffice ran"
        values = _xcu_values(seen['text'])
        for key, expected in _LOCKED_VALUES.items():
            assert values.get(key) == expected, (key, values)

    def test_profile_is_rewritten_if_libreoffice_changed_it(self, tmp_path, monkeypatch):
        # A conversion run must never be the one that finds the lock removed.
        monkeypatch.setattr(document_converter, '_TEMP_ROOT', tmp_path / 'root')
        (tmp_path / 'root').mkdir()
        xcu = document_converter._profile_dir() / 'user' / 'registrymodifications.xcu'
        xcu.write_text('<oor:items/>', encoding='utf-8')
        document_converter._profile_dir()
        assert _xcu_values(xcu.read_text(encoding='utf-8')) == _LOCKED_VALUES


class _RecordingHandler(http.server.BaseHTTPRequestHandler):
    requests = []

    def do_GET(self):
        type(self).requests.append(self.path)
        self.send_response(404)
        self.end_headers()

    do_HEAD = do_GET

    def log_message(self, *args):
        pass


def _hostile_docx(path, base_url):
    """
    A synthetic .docx whose picture is LINKED (r:link, TargetMode="External")
    to base_url, and whose settings attach a remote template. Opening it with
    links allowed makes LibreOffice request those URLs.
    """
    import io
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from PIL import Image

    png = io.BytesIO()
    Image.new('RGB', (4, 4), 'white').save(png, 'PNG')
    png.seek(0)

    document = docx.Document()
    document.add_paragraph('Student: Joe Bloggs, Dummy School.')
    run = document.add_paragraph().add_run()
    run.add_picture(png)
    blip = run._r.xpath('.//a:blip')[0]
    rid = document.part.relate_to(f'{base_url}/linked-image.png', RT.IMAGE, is_external=True)
    r_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    del blip.attrib[f'{{{r_ns}}}embed']
    blip.set(f'{{{r_ns}}}link', rid)
    document.save(str(path))

    # Remote attached template in settings.xml (a well-known fetch vector).
    with zipfile.ZipFile(path) as zin:
        items = {n: zin.read(n) for n in zin.namelist()}
    settings = items['word/settings.xml'].decode('utf-8')
    settings = settings.replace(
        '<w:zoom ', '<w:attachedTemplate r:id="rIdRemoteTpl"/><w:zoom ', 1)
    if 'xmlns:r=' not in settings.split('>', 2)[1]:
        settings = settings.replace(
            '<w:settings ', f'<w:settings xmlns:r="{r_ns}" ', 1)
    items['word/settings.xml'] = settings.encode('utf-8')
    items['word/_rels/settings.xml.rels'] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdRemoteTpl" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" '
        f'Target="{base_url}/remote-template.dotx" TargetMode="External"/>'
        '</Relationships>').encode('utf-8')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)


@pytest.mark.skipif(resolve_libreoffice() is None, reason="LibreOffice not installed")
def test_no_remote_fetch_from_linked_image(tmp_path):
    """
    Real LibreOffice, real conversion. The "remote" URLs point at a listener
    on 127.0.0.1, so a fetch is recorded here and nothing leaves the machine.
    """
    _RecordingHandler.requests = []
    server = http.server.HTTPServer(('127.0.0.1', 0), _RecordingHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        base_url = f'http://127.0.0.1:{server.server_address[1]}'
        docx_path = tmp_path / 'Joe Bloggs report.docx'
        _hostile_docx(docx_path, base_url)
        out = tmp_path / 'out'
        out.mkdir()
        ok, message, pdf = DocumentConverter().convert_to_pdf(docx_path, out)
    finally:
        server.shutdown()
        server.server_close()
    assert _RecordingHandler.requests == []
    # Locking the profile must not stop ordinary conversion.
    assert ok, message
    import fitz
    with fitz.open(str(pdf)) as doc:
        assert 'Joe Bloggs' in doc[0].get_text()


_FODT = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:xlink="http://www.w3.org/1999/xlink" office:version="1.3" '
    'office:mimetype="application/vnd.oasis.opendocument.text">'
    '<office:body><office:text>{body}</office:text></office:body></office:document>'
)


@pytest.mark.skipif(resolve_libreoffice() is None, reason="LibreOffice not installed")
def test_writer_link_never_does_not_update_linked_section(tmp_path, monkeypatch):
    """
    A Writer section linked to a LOCAL file (no network) whose cached text
    differs from the source. With the locked profile the conversion must keep
    the cached text, i.e. not follow the link. On LibreOffice 26.8 no Writer
    Update/Link value updates the section in headless conversion, so this is
    a guard against a version that starts to, not proof of the 0 = never
    mapping (rule 82).
    """
    monkeypatch.setattr(document_converter, '_TEMP_ROOT', tmp_path / 'root')
    (tmp_path / 'root').mkdir()
    source = tmp_path / 'source.fodt'
    source.write_text(_FODT.format(body='<text:p>FRESHSOURCETEXT</text:p>'), encoding='utf-8')
    main = tmp_path / 'main.fodt'
    main.write_text(_FODT.format(body=(
        '<text:p>Dummy School</text:p>'
        '<text:section text:name="S1">'
        f'<text:section-source xlink:href="{source.as_uri()}" xlink:type="simple"/>'
        '<text:p>CACHEDPLACEHOLDER</text:p></text:section>')), encoding='utf-8')
    out = tmp_path / 'out'
    out.mkdir()
    result = subprocess.run(
        [resolve_libreoffice(),
         f'-env:UserInstallation={document_converter._profile_dir().as_uri()}',
         '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(main)],
        capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stderr
    import fitz
    with fitz.open(str(out / 'main.pdf')) as doc:
        text = doc[0].get_text()
    assert 'CACHEDPLACEHOLDER' in text
    assert 'FRESHSOURCETEXT' not in text
